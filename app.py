import io, os, sqlite3
from datetime import datetime

from flask import (
    Flask,
    render_template,
    request,
    redirect,
    url_for,
    flash,
    send_file,
)
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from dotenv import load_dotenv


load_dotenv()

app = Flask(__name__)

app.config["SECRET_KEY"] = os.getenv(
    "FLASK_SECRET_KEY"
)

DATABASE_PATH = os.getenv(
    "DATABASE_PATH",
    "database/credit_debt.db"
)

os.makedirs(os.path.dirname(DATABASE_PATH), exist_ok=True)


# --------------------------------------------------
# DATABASE
# --------------------------------------------------

def get_db():
    conn = sqlite3.connect(DATABASE_PATH)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys = ON")
    return conn


def init_db():
    conn = get_db()

    conn.executescript("""
        CREATE TABLE IF NOT EXISTS debts (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            creditor TEXT NOT NULL,
            description TEXT,
            original_amount REAL NOT NULL DEFAULT 0,
            remaining_amount REAL NOT NULL DEFAULT 0,
            interest_rate REAL NOT NULL DEFAULT 0,
            minimum_payment REAL NOT NULL DEFAULT 0,
            due_date TEXT,
            status TEXT NOT NULL DEFAULT 'active',
            archived INTEGER NOT NULL DEFAULT 0,
            created_at TEXT NOT NULL
        );

        CREATE TABLE IF NOT EXISTS payments (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            debt_id INTEGER NOT NULL,
            amount REAL NOT NULL,
            created_at TEXT NOT NULL,
            FOREIGN KEY (debt_id)
                REFERENCES debts(id)
                ON DELETE CASCADE
        );

        CREATE TABLE IF NOT EXISTS savings (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT NOT NULL,
            target_amount REAL NOT NULL,
            current_amount REAL NOT NULL DEFAULT 0,
            contribution_amount REAL NOT NULL,
            archived INTEGER NOT NULL DEFAULT 0,
            created_at TEXT NOT NULL,
            updated_at TEXT NOT NULL
        );

        CREATE TABLE IF NOT EXISTS savings_transactions (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            saving_id INTEGER NOT NULL,
            amount REAL NOT NULL,
            created_at TEXT NOT NULL,
            FOREIGN KEY (saving_id)
                REFERENCES savings(id)
                ON DELETE CASCADE
        );
    """)

    # Safe migration for existing databases
    for table in ("debts", "savings"):
        cols = [r[1] for r in conn.execute(f"PRAGMA table_info({table})")]
        if "archived" not in cols:
            conn.execute(
                f"ALTER TABLE {table} "
                f"ADD COLUMN archived INTEGER NOT NULL DEFAULT 0"
            )

    # Backfill: mark already-paid debts and completed savings as archived
    conn.execute("""
        UPDATE debts
        SET archived = 1
        WHERE status = 'paid' AND archived = 0
    """)

    conn.execute("""
        UPDATE savings
        SET archived = 1
        WHERE current_amount >= target_amount
          AND archived = 0
    """)

    conn.commit()
    conn.close()


init_db()


# --------------------------------------------------
# HELPERS
# --------------------------------------------------

def now():
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def peso(value):
    return f"₱{value:,.2f}"


app.jinja_env.filters["peso"] = peso


# --------------------------------------------------
# DASHBOARD
# --------------------------------------------------

@app.route("/")
def index():
    conn = get_db()

    debts = conn.execute("""
        SELECT *
        FROM debts
        WHERE archived = 0
        ORDER BY
            CASE WHEN status = 'active' THEN 0 ELSE 1 END,
            due_date ASC,
            id DESC
    """).fetchall()

    savings = conn.execute("""
        SELECT *
        FROM savings
        WHERE archived = 0
        ORDER BY id DESC
    """).fetchall()

    archived_debts = conn.execute("""
        SELECT *
        FROM debts
        WHERE archived = 1
        ORDER BY id DESC
    """).fetchall()

    archived_savings = conn.execute("""
        SELECT *
        FROM savings
        WHERE archived = 1
        ORDER BY id DESC
    """).fetchall()

    total_debt = conn.execute("""
        SELECT COALESCE(SUM(remaining_amount), 0)
        FROM debts
        WHERE status = 'active' AND archived = 0
    """).fetchone()[0]

    total_savings = conn.execute("""
        SELECT COALESCE(SUM(current_amount), 0)
        FROM savings
        WHERE archived = 0
    """).fetchone()[0]

    conn.close()

    net_position = total_savings - total_debt

    return render_template(
        "index.html",
        debts=debts,
        savings=savings,
        archived_debts=archived_debts,
        archived_savings=archived_savings,
        total_debt=total_debt,
        total_savings=total_savings,
        net_position=net_position,
    )


# --------------------------------------------------
# ADD DEBT
# --------------------------------------------------

@app.route("/add-debt", methods=["GET", "POST"])
def add_debt():
    if request.method == "POST":
        creditor = request.form.get("creditor", "").strip()
        description = request.form.get("description", "").strip()

        try:
            original_amount = float(
                request.form.get("original_amount", 0)
            )

            interest_rate = float(
                request.form.get("interest_rate", 0)
            )

            minimum_payment = float(
                request.form.get("minimum_payment", 0)
            )

        except ValueError:
            flash("Please enter valid amounts.", "error")
            return redirect(url_for("add_debt"))

        due_date = request.form.get("due_date") or None

        if not creditor:
            flash("Creditor is required.", "error")
            return redirect(url_for("add_debt"))

        if original_amount < 0:
            flash("Amount cannot be negative.", "error")
            return redirect(url_for("add_debt"))

        conn = get_db()

        conn.execute("""
            INSERT INTO debts (
                creditor,
                description,
                original_amount,
                remaining_amount,
                interest_rate,
                minimum_payment,
                due_date,
                status,
                archived,
                created_at
            )
            VALUES (?, ?, ?, ?, ?, ?, ?, 'active', 0, ?)
        """, (
            creditor,
            description,
            original_amount,
            original_amount,
            interest_rate,
            minimum_payment,
            due_date,
            now()
        ))

        conn.commit()
        conn.close()

        flash("Debt added successfully.", "success")

        return redirect(url_for("index"))

    return render_template("add_debt.html")


# --------------------------------------------------
# ADD SAVINGS
# --------------------------------------------------

@app.route("/add-savings", methods=["GET", "POST"])
def add_savings():
    if request.method == "POST":
        name = request.form.get("name", "").strip()

        try:
            target_amount = float(
                request.form.get("target_amount", 0)
            )

            current_amount = float(
                request.form.get("current_amount", 0)
            )

            contribution_amount = float(
                request.form.get("contribution_amount", 0)
            )

        except ValueError:
            flash("Please enter valid amounts.", "error")
            return redirect(url_for("add_savings"))

        if not name:
            flash("Savings name is required.", "error")
            return redirect(url_for("add_savings"))

        if target_amount <= 0:
            flash("Target amount must be greater than zero.", "error")
            return redirect(url_for("add_savings"))

        if current_amount < 0:
            flash("Current amount cannot be negative.", "error")
            return redirect(url_for("add_savings"))

        if contribution_amount <= 0:
            flash(
                "Contribution amount must be greater than zero.",
                "error"
            )
            return redirect(url_for("add_savings"))

        current_amount = min(
            current_amount,
            target_amount
        )

        archived = 1 if current_amount >= target_amount else 0

        conn = get_db()

        cursor = conn.execute("""
            INSERT INTO savings (
                name,
                target_amount,
                current_amount,
                contribution_amount,
                archived,
                created_at,
                updated_at
            )
            VALUES (?, ?, ?, ?, ?, ?, ?)
        """, (
            name,
            target_amount,
            current_amount,
            contribution_amount,
            archived,
            now(),
            now()
        ))

        saving_id = cursor.lastrowid

        if current_amount > 0:
            conn.execute("""
                INSERT INTO savings_transactions (
                    saving_id,
                    amount,
                    created_at
                )
                VALUES (?, ?, ?)
            """, (
                saving_id,
                current_amount,
                now()
            ))

        conn.commit()
        conn.close()

        flash("Savings goal added successfully.", "success")

        return redirect(url_for("index"))

    return render_template("add_savings.html")


# --------------------------------------------------
# MARK SAVINGS AS SAVED
# --------------------------------------------------

@app.route(
    "/savings/<int:saving_id>/save",
    methods=["POST"]
)
def mark_savings_saved(saving_id):
    conn = get_db()

    saving = conn.execute("""
        SELECT *
        FROM savings
        WHERE id = ?
    """, (saving_id,)).fetchone()

    if not saving:
        conn.close()
        flash("Savings goal not found.", "error")
        return redirect(url_for("index"))

    current = float(saving["current_amount"])
    target = float(saving["target_amount"])
    contribution = float(saving["contribution_amount"])

    if current >= target:
        conn.close()
        flash("This savings goal is already complete.", "error")
        return redirect(url_for("index"))

    actual_amount = min(
        contribution,
        target - current
    )

    new_amount = current + actual_amount
    archived = 1 if new_amount >= target else 0

    conn.execute("""
        UPDATE savings
        SET current_amount = ?,
            updated_at = ?,
            archived = ?
        WHERE id = ?
    """, (
        new_amount,
        now(),
        archived,
        saving_id
    ))

    conn.execute("""
        INSERT INTO savings_transactions (
            saving_id,
            amount,
            created_at
        )
        VALUES (?, ?, ?)
    """, (
        saving_id,
        actual_amount,
        now()
    ))

    conn.commit()
    conn.close()

    flash(
        f"{peso(actual_amount)} added to {saving['name']}.",
        "success"
    )

    return redirect(url_for("index"))


# --------------------------------------------------
# ARCHIVE / UNARCHIVE SAVINGS
# --------------------------------------------------

@app.route(
    "/savings/<int:saving_id>/archive",
    methods=["POST"]
)
def archive_savings(saving_id):
    conn = get_db()
    conn.execute(
        "UPDATE savings SET archived = 1 WHERE id = ?",
        (saving_id,)
    )
    conn.commit()
    conn.close()
    flash("Savings goal archived.", "success")
    return redirect(url_for("index"))


@app.route(
    "/savings/<int:saving_id>/unarchive",
    methods=["POST"]
)
def unarchive_savings(saving_id):
    conn = get_db()
    conn.execute(
        "UPDATE savings SET archived = 0 WHERE id = ?",
        (saving_id,)
    )
    conn.commit()
    conn.close()
    flash("Savings goal restored.", "success")
    return redirect(url_for("index"))


# --------------------------------------------------
# DELETE SAVINGS
# --------------------------------------------------

@app.route(
    "/delete-savings/<int:saving_id>",
    methods=["POST"]
)
def delete_savings(saving_id):
    conn = get_db()

    conn.execute("""
        DELETE FROM savings
        WHERE id = ?
    """, (saving_id,))

    conn.commit()
    conn.close()

    flash("Savings goal deleted.", "success")

    return redirect(url_for("index"))


# --------------------------------------------------
# DEBT DETAILS
# --------------------------------------------------

@app.route("/debt/<int:debt_id>")
def debt_detail(debt_id):
    conn = get_db()

    debt = conn.execute("""
        SELECT *
        FROM debts
        WHERE id = ?
    """, (debt_id,)).fetchone()

    if not debt:
        conn.close()
        flash("Debt not found.", "error")
        return redirect(url_for("index"))

    payments = conn.execute("""
        SELECT *
        FROM payments
        WHERE debt_id = ?
        ORDER BY created_at DESC
    """, (debt_id,)).fetchall()

    conn.close()

    return render_template(
        "debt.html",
        debt=debt,
        payments=payments
    )


# --------------------------------------------------
# PAYMENT
# --------------------------------------------------

@app.route(
    "/payment/<int:debt_id>",
    methods=["POST"]
)
def payment(debt_id):
    try:
        amount = float(
            request.form.get("amount", 0)
        )
    except ValueError:
        flash("Invalid payment amount.", "error")
        return redirect(
            url_for("debt_detail", debt_id=debt_id)
        )

    if amount <= 0:
        flash("Payment must be greater than zero.", "error")
        return redirect(
            url_for("debt_detail", debt_id=debt_id)
        )

    conn = get_db()

    debt = conn.execute("""
        SELECT *
        FROM debts
        WHERE id = ?
    """, (debt_id,)).fetchone()

    if not debt:
        conn.close()
        flash("Debt not found.", "error")
        return redirect(url_for("index"))

    remaining = float(debt["remaining_amount"])

    if remaining <= 0:
        conn.close()
        flash("This debt is already paid.", "error")
        return redirect(
            url_for("debt_detail", debt_id=debt_id)
        )

    actual_payment = min(amount, remaining)
    new_remaining = remaining - actual_payment

    status = "paid" if new_remaining <= 0 else "active"
    archived = 1 if new_remaining <= 0 else 0

    conn.execute("""
        UPDATE debts
        SET remaining_amount = ?,
            status = ?,
            archived = ?
        WHERE id = ?
    """, (
        new_remaining,
        status,
        archived,
        debt_id
    ))

    conn.execute("""
        INSERT INTO payments (
            debt_id,
            amount,
            created_at
        )
        VALUES (?, ?, ?)
    """, (
        debt_id,
        actual_payment,
        now()
    ))

    conn.commit()
    conn.close()

    flash(
        f"Payment of {peso(actual_payment)} recorded.",
        "success"
    )

    return redirect(
        url_for("debt_detail", debt_id=debt_id)
    )


# --------------------------------------------------
# ARCHIVE / UNARCHIVE DEBT
# --------------------------------------------------

@app.route(
    "/debt/<int:debt_id>/archive",
    methods=["POST"]
)
def archive_debt(debt_id):
    conn = get_db()
    conn.execute(
        "UPDATE debts SET archived = 1 WHERE id = ?",
        (debt_id,)
    )
    conn.commit()
    conn.close()
    flash("Debt archived.", "success")
    return redirect(url_for("index"))


@app.route(
    "/debt/<int:debt_id>/unarchive",
    methods=["POST"]
)
def unarchive_debt(debt_id):
    conn = get_db()
    conn.execute(
        "UPDATE debts SET archived = 0 WHERE id = ?",
        (debt_id,)
    )
    conn.commit()
    conn.close()
    flash("Debt restored.", "success")
    return redirect(url_for("index"))


# --------------------------------------------------
# DELETE DEBT
# --------------------------------------------------

@app.route(
    "/delete-debt/<int:debt_id>",
    methods=["POST"]
)
def delete_debt(debt_id):
    conn = get_db()

    conn.execute("""
        DELETE FROM debts
        WHERE id = ?
    """, (debt_id,))

    conn.commit()
    conn.close()

    flash("Debt deleted.", "success")

    return redirect(url_for("index"))


# --------------------------------------------------
# EXPORT EXCEL
# --------------------------------------------------

@app.route("/export-excel")
def export_excel():
    conn = get_db()

    debts = conn.execute("""
        SELECT *
        FROM debts
        ORDER BY id
    """).fetchall()

    savings = conn.execute("""
        SELECT *
        FROM savings
        ORDER BY id
    """).fetchall()

    payments = conn.execute("""
        SELECT
            payments.id,
            debts.creditor,
            payments.amount,
            payments.created_at
        FROM payments
        JOIN debts
            ON debts.id = payments.debt_id
        ORDER BY payments.created_at DESC
    """).fetchall()

    savings_transactions = conn.execute("""
        SELECT
            savings_transactions.id,
            savings.name,
            savings_transactions.amount,
            savings_transactions.created_at
        FROM savings_transactions
        JOIN savings
            ON savings.id = savings_transactions.saving_id
        ORDER BY savings_transactions.created_at DESC
    """).fetchall()

    total_debt = conn.execute("""
        SELECT COALESCE(SUM(remaining_amount), 0)
        FROM debts
        WHERE status = 'active' AND archived = 0
    """).fetchone()[0]

    total_savings = conn.execute("""
        SELECT COALESCE(SUM(current_amount), 0)
        FROM savings
        WHERE archived = 0
    """).fetchone()[0]

    conn.close()

    # ----------------------------------------------
    # CREATE WORKBOOK
    # ----------------------------------------------

    workbook = Workbook()

    summary = workbook.active
    summary.title = "Summary"

    summary.append(["Credit & Debt"])
    summary.append([])
    summary.append(["Total Debt", total_debt])
    summary.append(["Total Savings", total_savings])
    summary.append([
        "Net Position",
        total_savings - total_debt
    ])

    # ----------------------------------------------
    # SAVINGS
    # ----------------------------------------------

    ws = workbook.create_sheet("Savings")

    ws.append([
        "ID",
        "Name",
        "Target Amount",
        "Current Amount",
        "Remaining",
        "Contribution Amount",
        "Progress",
        "Archived",
        "Created At",
        "Updated At"
    ])

    for saving in savings:
        target = float(saving["target_amount"])
        current = float(saving["current_amount"])

        remaining = max(target - current, 0)

        progress = (
            current / target
            if target > 0
            else 0
        )

        ws.append([
            saving["id"],
            saving["name"],
            target,
            current,
            remaining,
            saving["contribution_amount"],
            progress,
            "Yes" if saving["archived"] else "No",
            saving["created_at"],
            saving["updated_at"]
        ])

    # ----------------------------------------------
    # DEBTS
    # ----------------------------------------------

    ws = workbook.create_sheet("Debts")

    ws.append([
        "ID",
        "Creditor",
        "Description",
        "Original Amount",
        "Remaining Amount",
        "Interest Rate",
        "Minimum Payment",
        "Due Date",
        "Status",
        "Archived",
        "Created At"
    ])

    for debt in debts:
        ws.append([
            debt["id"],
            debt["creditor"],
            debt["description"],
            debt["original_amount"],
            debt["remaining_amount"],
            debt["interest_rate"],
            debt["minimum_payment"],
            debt["due_date"],
            debt["status"],
            "Yes" if debt["archived"] else "No",
            debt["created_at"]
        ])

    # ----------------------------------------------
    # PAYMENTS
    # ----------------------------------------------

    ws = workbook.create_sheet("Payments")

    ws.append([
        "ID",
        "Creditor",
        "Amount",
        "Date"
    ])

    for payment_row in payments:
        ws.append([
            payment_row["id"],
            payment_row["creditor"],
            payment_row["amount"],
            payment_row["created_at"]
        ])

    # ----------------------------------------------
    # SAVINGS HISTORY
    # ----------------------------------------------

    ws = workbook.create_sheet("Savings History")

    ws.append([
        "ID",
        "Savings Goal",
        "Amount",
        "Date"
    ])

    for transaction in savings_transactions:
        ws.append([
            transaction["id"],
            transaction["name"],
            transaction["amount"],
            transaction["created_at"]
        ])

    # ----------------------------------------------
    # STYLE WORKBOOK
    # ----------------------------------------------

    for worksheet in workbook.worksheets:
        worksheet.freeze_panes = "A2"

        for cell in worksheet[1]:
            cell.font = Font(bold=True)

        for column in worksheet.columns:
            max_length = 0
            column_letter = column[0].column_letter

            for cell in column:
                try:
                    value_length = len(str(cell.value))
                    max_length = max(
                        max_length,
                        value_length
                    )
                except Exception:
                    pass

            worksheet.column_dimensions[
                column_letter
            ].width = min(max_length + 2, 40)

    # Percentage formatting
    savings_sheet = workbook["Savings"]

    for row in range(
        2,
        savings_sheet.max_row + 1
    ):
        savings_sheet.cell(
            row=row,
            column=7
        ).number_format = "0.00%"

    # Currency formatting
    for sheet_name, columns in {
        "Summary": [2],
        "Savings": [3, 4, 5, 6],
        "Debts": [4, 5, 7],
        "Payments": [3],
        "Savings History": [3],
    }.items():

        worksheet = workbook[sheet_name]

        for column in columns:
            for row in range(
                2,
                worksheet.max_row + 1
            ):
                worksheet.cell(
                    row=row,
                    column=column
                ).number_format = '₱#,##0.00'

    # ----------------------------------------------
    # SEND FILE
    # ----------------------------------------------

    output = io.BytesIO()

    workbook.save(output)
    output.seek(0)

    filename = (
        f"credit_debt_"
        f"{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    )

    return send_file(
        output,
        as_attachment=True,
        download_name=filename,
        mimetype=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        )
    )


# --------------------------------------------------
# RUN
# --------------------------------------------------

if __name__ == "__main__":
    app.run(
        host="0.0.0.0",
        port=int(os.getenv("PORT", 5000)),
        debug=os.getenv("DEBUG", "true").lower() == "true"
    )